"""Build the five-slide summary for a marketing manager from the generated results.

    python slides/build_slides.py      # after `python -m email_uplift.report`

Every number is read from reports/results/ (summary.json, segment_lifts.csv), so the deck never disagrees with
the analysis.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "reports" / "figures"
OUT = ROOT / "slides" / "email-campaign-summary.pptx"
INK, INK_2, ACCENT = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x2A, 0x78, 0xD6)
W, H = Inches(13.333), Inches(7.5)


def _text(slide, left, top, width, height, paragraphs, size=16, color=INK, bold_first=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, text in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold_first and i == 0
        p.space_after = Pt(10)
    return box


def _slide(prs, title, kicker=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    if kicker:
        _text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.4), [kicker], 12, ACCENT)
    _text(s, Inches(0.6), Inches(0.65), Inches(12.1), Inches(1.1), [title], 26, INK, True)
    return s


def _picture(slide, path, left, top, width=None, height=None):
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def _usd(v):
    return f"{'-' if v < 0 else ''}${abs(v) / 1000:,.1f}k"


def build(summary: dict, segments: pd.DataFrame) -> Presentation:
    lift, pol, tgt = summary["lift"], summary["policies"], summary["targeting"]
    m_conv, w_conv = lift["mens_vs_control_conversion"], lift["womens_vs_control_conversion"]
    m_sp, w_sp = lift["mens_vs_control_spend"], lift["womens_vs_control_spend"]
    mw_conv, mw_sp = lift["mens_vs_womens_conversion"], lift["mens_vs_womens_spend"]
    everyone = pol["Everyone, men's email (full budget)"]["total_per_customer"]
    nobody = pol["No email"]["total_per_customer"]
    rand_half = pol["Random half, men's email"]["total_per_customer"]
    spenders = pol["Top past spenders half, men's email"]
    tw = tgt["womens"]
    worst_p = max(lift[f"{a}_vs_control_{m}"]["p_holm"] for a in ("mens", "womens")
                  for m in ("visit", "conversion", "spend"))

    def visit_pp(email, level):
        r = segments[(segments["metric"] == "visit") & (segments["dimension"] == "Past purchase")
                     & (segments["treatment"] == email) & (segments["level"] == level)]
        return 100 * float(r["abs_lift"].iloc[0])

    visit_men_only = visit_pp("womens", "Men's only")
    visit_women_only = visit_pp("womens", "Women's only")

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1. Bottom line
    s = _slide(prs, "Email the whole list with the men's email. If you can only reach half, "
                    "the men's email still wins, and how you pick that half matters little.",
               "Email test · 64,000 customers · results over two weeks")
    _text(s, Inches(0.6), Inches(2.2), Inches(12), Inches(4.8), [
        f"• Both emails worked. The men's email more than doubled the purchase rate "
        f"({100 * m_conv['control_rate']:.2f}% → {100 * m_conv['treatment_rate']:.2f}%) and added "
        f"${m_sp['abs_lift']:.2f} of spend per customer.",
        f"• The men's email beat the women's email on visits, purchases and spend.",
        f"• Emailing everyone with the men's email: about {_usd(everyone - nobody)} more spend in two "
        f"weeks than sending nothing. A random half gets about half of that ({_usd(rand_half - nobody)}).",
        f"• Its effect is broad. No customer group was left out, so a targeting model could not "
        f"pick a half that clearly beats a random half.",
        f"• The women's email is different. Its effect is concentrated, so if it is sent, aim it: the "
        f"top-ranked half delivers {tw['conversions_share_of_full_campaign']:.0%} of its extra "
        f"purchases.",
    ], 18)

    # 2. Did it work?
    s = _slide(prs, "Both emails caused more visits, purchases and spend; it is not chance",
               "1 · Did it work?")
    _picture(s, FIG / "lift_overview.png", Inches(0.6), Inches(1.8), width=Inches(12.1))
    _text(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.6), [
        f"Customers were assigned at random, and the three groups looked alike beforehand "
        f"(largest difference {summary['balance_max_abs_smd']:.3f} standard deviations), so the "
        f"gaps are caused by the emails. Every lift holds at p ≤ {worst_p:.3f} after correcting for "
        f"running nine tests. Men's email: +{100 * m_conv['abs_lift']:.2f} pp purchase rate "
        f"(+{m_conv['rel_lift']:.0%}); women's: +{100 * w_conv['abs_lift']:.2f} pp "
        f"(+{w_conv['rel_lift']:.0%}).",
    ], 15, INK_2)

    # 3. Which won, and where
    s = _slide(prs, "The men's email won. The women's email did nothing measurable for men's-only "
                    "buyers or returning customers", "2 · Which email, and for whom?")
    pic = _picture(s, FIG / "segment_conversion_lift.png", Inches(7.2), Inches(1.75), height=Inches(5.6))
    _text(s, Inches(0.6), Inches(2.0), Inches(6.3), Inches(5), [
        f"• Head to head, the men's email had +{100 * mw_conv['abs_lift']:.2f} pp higher purchase rate "
        f"and +${mw_sp['abs_lift']:.2f} more spend per customer than the women's email.",
        "• Men's email: positive in every segment we cut, including women's-only buyers.",
        "• Women's email: clear lift for women's buyers, new customers and $500+ spenders; no "
        "detectable lift for men's-only buyers, returning customers, or $200–500 spenders.",
        "• Caution: most segment differences are within noise once you account for testing "
        "many segments. The one solid difference: the women's email barely moved site visits from "
        f"men's-only buyers (+{visit_men_only:.1f} pp vs "
        f"+{visit_women_only:.1f} pp for women's-only buyers), while the men's "
        "email drew both groups equally and buyers of both categories most.",
    ], 15)

    # 4. Targeting
    s = _slide(prs, "An uplift model ranks who the email changes, not who buys anyway",
               "3 · Who should get it?")
    _picture(s, FIG / "uplift_curves.png", Inches(1.5), Inches(1.65), height=Inches(4.3))
    _text(s, Inches(0.6), Inches(6.05), Inches(12.1), Inches(1.3), [
        f"Women's email: the model's top half delivers {_usd(tw['spend_incremental_model'])} of extra "
        f"spend vs {_usd(tw['spend_incremental_random'])} for a random half, and "
        f"{tw['conversions_share_of_full_campaign']:.0%} of the extra purchases. Men's email: the "
        f"ranked half does no better than a random half, because the email lifts nearly everyone.",
    ], 15, INK_2)

    # 5. Recommendation
    s = _slide(prs, "Recommendation", "4 · What to do next")
    _picture(s, FIG / "half_budget_policies.png", Inches(6.6), Inches(1.7), width=Inches(6.4))
    _text(s, Inches(0.6), Inches(1.8), Inches(5.8), Inches(5.4), [
        "1. Default: send the men's email to everyone. Sending costs little, and every segment "
        "responded.",
        f"2. Half-list budget: send the men's email. A random half earns {_usd(rand_half)}. "
        f"Picking the top past spenders scored {_usd(spenders['total_vs_reference'])} more, "
        f"but that could be chance (95% CI {_usd(spenders['total_vs_reference_ci_low'])} to "
        f"{_usd(spenders['total_vs_reference_ci_high'])}).",
        "3. If the women's email must go out, send it only to the model's top half, mainly "
        "women's buyers and new customers.",
        "4. Next test: next send, run past-spend targeting against random selection, and give "
        "both groups the men's email, to confirm the extra revenue before rolling it out.",
    ], 15)
    return prs


def main() -> None:
    results = ROOT / "reports" / "results"
    summary = json.loads((results / "summary.json").read_text())
    build(summary, pd.read_csv(results / "segment_lifts.csv")).save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
