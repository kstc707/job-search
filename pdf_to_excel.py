"""
Convert header-less tabular data in a PDF into an Excel file with headers.

Running from VS Code
--------------------
  1. Put this script and your PDF in the same folder.
  2. Edit the SETTINGS block below (headers + rows to skip).
  3. Press the Run button (top-right) or F5. The Excel file is saved
     next to the PDF.

Command-line usage examples (optional; overrides SETTINGS)
--------------
  # Headers given on the command line, skip the first 2 rows
  python pdf_to_excel.py data.pdf -o output.xlsx \
      --headers "ID,Name,Date,Amount" --skip-first 2

  # Also drop specific rows (0-based, counted after extraction) and any row
  # containing a pattern such as page footers
  python pdf_to_excel.py data.pdf --headers "ID,Name,Date,Amount" \
      --skip-rows 5,17 --skip-pattern "Page \\d+ of \\d+" --skip-pattern "^Total"

  # Only use pages 2-4
  python pdf_to_excel.py data.pdf --headers "A,B,C" --pages 2-4

Requires:  pip install pdfplumber pandas openpyxl
"""

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================ SETTINGS ============================
# Used when you press Run in VS Code (no command-line arguments).

# PDF file name in the same folder as this script.
# Leave as "" to automatically use the first .pdf found in the folder.
PDF_FILE = ""

# Name of the Excel file to create (saved in the same folder).
OUTPUT_FILE = "output.xlsx"

# Your column headers, in order.
HEADERS = ["ID", "Name", "Date", "Amount"]

# Number of rows to skip at the very top (titles, notes, etc.).
SKIP_FIRST = 0

# Specific row numbers to drop (see PREVIEW output; counting starts at 0).
SKIP_ROWS = []

# Drop any row whose text matches one of these regex patterns.
SKIP_PATTERNS = [r"Page \d+ of \d+"]

# Pages to read, e.g. "1,3-5". None = all pages.
PAGES = None

# "auto", "table" (PDF has grid lines) or "text" (space-aligned columns).
MODE = "auto"

# True = only print the extracted rows with their numbers (no Excel).
# Run once with True to decide SKIP_FIRST / SKIP_ROWS, then set to False.
PREVIEW = True
# ==================================================================


def parse_pages(spec, total):
    """'1,3,5-7' -> [0, 2, 4, 5, 6] (0-based page indexes)."""
    if not spec:
        return list(range(total))
    pages = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            start, end = part.split("-")
            pages.extend(range(int(start) - 1, int(end)))
        else:
            pages.append(int(part) - 1)
    return [p for p in pages if 0 <= p < total]


def extract_rows(pdf_path, page_spec=None, mode="auto"):
    """Pull rows out of the PDF.

    mode="table": use pdfplumber's table detection (works when the PDF has
                  ruled lines / clear columns).
    mode="text":  split each text line on 2+ spaces (works for plain,
                  space-aligned data).
    mode="auto":  try tables first, fall back to text per page.
    """
    rows = []
    with pdfplumber.open(pdf_path) as pdf:
        for idx in parse_pages(page_spec, len(pdf.pages)):
            page = pdf.pages[idx]
            page_rows = []

            if mode in ("auto", "table"):
                for table in page.extract_tables():
                    page_rows.extend(table)

            if not page_rows and mode in ("auto", "text"):
                text = page.extract_text(layout=True) or ""
                for line in text.splitlines():
                    if line.strip():
                        page_rows.append(re.split(r"\s{2,}", line.strip()))

            rows.extend(page_rows)
    return rows


def clean_rows(rows, skip_first=0, skip_rows=None, skip_patterns=None,
               drop_empty=True):
    """Apply all row-skipping rules."""
    skip_rows = set(skip_rows or [])
    patterns = [re.compile(p) for p in (skip_patterns or [])]

    cleaned = []
    for i, row in enumerate(rows):
        if i < skip_first or i in skip_rows:
            continue
        cells = ["" if c is None else str(c).replace("\n", " ").strip()
                 for c in row]
        if drop_empty and not any(cells):
            continue
        line = " ".join(cells)
        if any(p.search(line) for p in patterns):
            continue
        cleaned.append(cells)
    return cleaned


def fit_to_headers(rows, headers):
    """Pad short rows / merge overflow so every row has len(headers) cells."""
    n = len(headers)
    fixed = []
    for row in rows:
        if len(row) < n:
            row = row + [""] * (n - len(row))
        elif len(row) > n:
            # Join extra cells into the last column rather than losing data
            row = row[: n - 1] + [" ".join(row[n - 1:])]
        fixed.append(row)
    return fixed


def to_number(series):
    """Convert a column to numeric where every non-blank value is numeric."""
    stripped = series.str.replace(r"[,$%]", "", regex=True).str.strip()
    numeric = pd.to_numeric(stripped.replace("", pd.NA), errors="coerce")
    non_blank = stripped != ""
    if non_blank.any() and numeric[non_blank].notna().all():
        return numeric
    return series


def write_excel(df, out_path, sheet_name="Data"):
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        header_fill = PatternFill("solid", fgColor="D9E1F2")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        for col_idx, col in enumerate(df.columns, start=1):
            longest = max([len(str(col))] + [len(str(v)) for v in df[col]])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(longest + 2, 60)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser(description="PDF (no headers) -> Excel with headers")
    ap.add_argument("pdf", help="input PDF file")
    ap.add_argument("-o", "--output", default="output.xlsx", help="output .xlsx path")
    ap.add_argument("--headers",
                    help='comma-separated column names, e.g. "ID,Name,Date,Amount"')
    ap.add_argument("--skip-first", type=int, default=0,
                    help="number of leading rows to skip (titles, notes, etc.)")
    ap.add_argument("--skip-rows", default="",
                    help="comma-separated 0-based row numbers to drop, e.g. 5,17")
    ap.add_argument("--skip-pattern", action="append", default=[],
                    help="regex; drop any row whose text matches (repeatable)")
    ap.add_argument("--pages", default=None, help='pages to read, e.g. "1,3-5"')
    ap.add_argument("--mode", choices=["auto", "table", "text"], default="auto",
                    help="extraction method (default: auto)")
    ap.add_argument("--no-numeric", action="store_true",
                    help="keep every column as text")
    ap.add_argument("--preview", action="store_true",
                    help="print the raw extracted rows with their numbers and exit")
    if len(sys.argv) == 1:
        # Launched from VS Code's Run button: use the SETTINGS block.
        folder = Path(__file__).resolve().parent
        if PDF_FILE:
            pdf_path = folder / PDF_FILE
        else:
            pdfs = sorted(folder.glob("*.pdf"))
            if not pdfs:
                sys.exit(f"No PDF found in {folder}. Put your PDF there "
                         "or set PDF_FILE.")
            pdf_path = pdfs[0]
        if not pdf_path.exists():
            sys.exit(f"PDF not found: {pdf_path}")
        print(f"Reading {pdf_path.name}")
        argv = [str(pdf_path), "-o", str(folder / OUTPUT_FILE),
                "--headers", ",".join(HEADERS),
                "--skip-first", str(SKIP_FIRST),
                "--skip-rows", ",".join(map(str, SKIP_ROWS)),
                "--mode", MODE]
        for pat in SKIP_PATTERNS:
            argv += ["--skip-pattern", pat]
        if PAGES:
            argv += ["--pages", PAGES]
        if PREVIEW:
            argv.append("--preview")
        args = ap.parse_args(argv)
    else:
        args = ap.parse_args()

    raw = extract_rows(args.pdf, args.pages, args.mode)
    if not raw:
        sys.exit("No rows found. If the PDF is a scanned image it needs OCR first.")

    if args.preview:
        for i, r in enumerate(raw):
            print(f"{i:>4}: {r}")
        print("\nPreview only. Set PREVIEW = False (or drop --preview) "
              "to create the Excel file.")
        return

    if not args.headers:
        sys.exit("--headers is required (run with --preview first to see the rows).")
    headers = [h.strip() for h in args.headers.split(",")]
    skip_rows = [int(x) for x in args.skip_rows.split(",") if x.strip()]

    rows = clean_rows(raw, args.skip_first, skip_rows, args.skip_pattern)
    rows = fit_to_headers(rows, headers)

    df = pd.DataFrame(rows, columns=headers)
    if not args.no_numeric:
        for col in df.columns:
            df[col] = to_number(df[col])

    write_excel(df, args.output)
    print(f"Wrote {len(df)} rows x {len(headers)} columns to {args.output} "
          f"(extracted {len(raw)}, skipped {len(raw) - len(df)})")


if __name__ == "__main__":
    main()
